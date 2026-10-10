"""OMML DOCX regressions through preview, persisted ingestion and retrieval."""
import io
import uuid
from unittest.mock import AsyncMock

import docx
import fitz
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.core.config import settings
from app.main import app
from app.ml.mock_provider import MockEmbeddingProvider
from app.ml.reranker_mock import MockRerankerProvider
from app.models.document import Document
from app.schemas.retrieval import RetrievalRequest
from app.services import document_preview
from app.services.document_service import DocumentService
from app.services.ingestion.parsers.docx import DocxParser
from app.services.ingestion.parsers.docx_math import read_equation
from app.services.ingestion.pipeline import IngestionPipeline
from app.services.ingestion_service import IngestionService, _IN_MEMORY_CHUNKS
from app.services.retrieval_service import RetrievalService

FRACTIONS = ["4/25", "8/50", "7/20", "8/400", "36/900", "9/125"]
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
INTRO = "Bài 1: Viết các phân số sau dưới dạng phân số thập phân."


def math_xml(body):
    return parse_xml(f'<m:oMath {nsdecls("m")}>{body}</m:oMath>')


def run(text):
    return f"<m:r><m:t>{text}</m:t></m:r>"


def fraction(num, den):
    return f"<m:f><m:num>{run(num)}</m:num><m:den>{run(den)}</m:den></m:f>"


def save(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def equation_docx():
    """Real zipped DOCX with inline/block Word math and a table equation."""
    document = docx.Document()
    document.add_heading("Ôn tập phân số", level=1)
    paragraph = document.add_paragraph(INTRO + " ")
    for index, value in enumerate(FRACTIONS[:4]):
        num, den = value.split("/")
        paragraph._p.append(math_xml(fraction(num, den)))
        paragraph.add_run("; " if index < 3 else ".")
    block = document.add_paragraph("Tiếp theo: ")
    block._p.append(parse_xml(f'<m:oMathPara {nsdecls("m")}><m:oMath>{fraction("36", "900")}</m:oMath></m:oMathPara>'))
    block.add_run(" rồi đến bảng.")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Cuối cùng: ")
    cell.paragraphs[0]._p.append(math_xml(fraction("9", "125")))
    document.add_paragraph("Kết thúc bài tập.")
    return save(document)


def assert_equations(text):
    assert INTRO in text
    positions = [text.index(value) for value in FRACTIONS]
    assert positions == sorted(positions)
    assert text.index("Kết thúc bài tập.") > positions[-1]


def test_ordered_equations_in_paragraphs_blocks_tables_and_pipeline(equation_docx):
    text = DocxParser().parse(equation_docx, "phân số.docx")[0].text
    assert_equations(text)
    assert text.startswith("# Ôn tập phân số")
    assert "8/400.\n\nTiếp theo: 36/900 rồi đến bảng." in text
    chunks = IngestionPipeline().process(equation_docx, "phân số.docx", MIME)
    assert_equations("\n".join(chunk.content for chunk in chunks))
    assert all(chunk.page_number_start == chunk.page_number_end == 1 for chunk in chunks)


@pytest.mark.parametrize("body,expected", [
    (fraction("4", "25"), "4/25"),
    (f"<m:sSup><m:e>{run('x')}</m:e><m:sup>{run('2')}</m:sup></m:sSup>", "x^2"),
    (fraction("x+1", "y-2"), "(x+1)/(y-2)"),
    (f"<m:f><m:num>{fraction('1', '2')}</m:num><m:den>{run('3')}</m:den></m:f>", "(1/2)/3"),
    (run("2×3−1÷2=5"), "2*3-1/2=5"),
    (f"<m:sSup><m:e>{run('x+1')}</m:e><m:sup>{run('n+2')}</m:sup></m:sSup>", "(x+1)^(n+2)"),
    (f"<m:rad><m:e>{run('x')}</m:e></m:rad>", "sqrt(x)"),
])
def test_canonical_math(body, expected):
    equation = read_equation(math_xml(body))
    assert equation.text == expected
    assert equation.supported


@pytest.mark.parametrize("body,expected", [
    (f"<m:acc><m:e>{run('x+1')}</m:e></m:acc>", "[unsupported acc: x+1]"),
    ("<m:acc/>", "[unsupported equation]"),
    (f"<m:f><m:num>{run('4')}</m:num></m:f>", "4/?"),
])
def test_unsupported_or_incomplete_math_retains_content_without_crashing(body, expected, caplog):
    document = docx.Document()
    paragraph = document.add_paragraph("Trước: ")
    paragraph._p.append(math_xml(body))
    paragraph.add_run(" Sau.")
    text = DocxParser().parse(save(document), "unknown.docx")[0].text
    assert text == f"Trước: {expected} Sau."
    assert "docx_math_fallback node=" in caplog.text
    assert "Trước" not in caplog.text
    assert not read_equation(math_xml(body)).supported


@pytest.mark.asyncio
async def test_equations_reach_embedding_and_hybrid_retrieval(equation_docx, monkeypatch):
    user, workspace, document_id = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    document = Document(id=document_id, user_id=user, workspace_id=workspace,
                        original_filename="phân số.docx", mime_type=MIME, status="uploaded")
    monkeypatch.setattr(DocumentService, "get_document", AsyncMock(return_value=document))
    monkeypatch.setattr(DocumentService, "get_document_file", AsyncMock(return_value=(document, equation_docx)))
    monkeypatch.setattr(settings, "LOCAL_SHARED_MODELS", False)
    monkeypatch.setattr(settings, "AUTO_EMBED_AFTER_INGESTION", True)
    provider = MockEmbeddingProvider()
    embedded = []
    encode = provider.encode_batch
    def record(texts, **kwargs):
        embedded.extend(texts)
        return encode(texts, **kwargs)
    monkeypatch.setattr(provider, "encode_batch", record)
    monkeypatch.setattr("app.services.embedding_service.get_embedding_provider", lambda: provider)
    try:
        assert await IngestionService.process_document(None, document_id, user)
        assert document.embedding_status == "completed"
        assert_equations("\n".join(embedded))
        chunks = await IngestionService.list_document_chunks(None, document_id, user)
        assert_equations("\n".join(chunk.content for chunk in chunks))
        response = await RetrievalService.retrieve(
            None, workspace, user,
            RetrievalRequest(query="Bài 1 phân số 4/25", routing_mode="always_quality", relevance_threshold=0),
            embedding_provider=provider, reranker_provider=MockRerankerProvider(),
        )
        assert response.results
        assert_equations("\n".join(chunk.content for chunk in response.results))
        assert all(chunk.document_id == document_id for chunk in response.results)
    finally:
        _IN_MEMORY_CHUNKS.pop(document_id, None)


def pdf_with_text(text):
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((40, 50), text)
        return pdf.tobytes()


@pytest.mark.parametrize("native", ["missing", "stacked", "invalid", "unavailable", "preserved"])
def test_preview_keeps_readable_equations_and_owner_isolation(equation_docx, monkeypatch, native):
    user = uuid.uuid4()
    headers = {"Authorization": f"Bearer test-token:{user}"}
    client = TestClient(app)
    workspace = client.post("/api/v1/workspaces", headers=headers, json={"name": "Math preview"}).json()["id"]
    uploaded = client.post(f"/api/v1/workspaces/{workspace}/documents", headers=headers,
                           files={"file": ("phân số.docx", equation_docx, MIME)})
    assert uploaded.status_code == 201
    document_id = uuid.UUID(uploaded.json()["id"])
    output = pdf_with_text(" ".join(FRACTIONS) if native == "preserved" else
                           "\n".join(FRACTIONS).replace("/", "\n") if native == "stacked" else "Missing mathematics")
    def convert(*args):
        if native == "unavailable":
            raise RuntimeError("PREVIEW_CONVERTER_UNAVAILABLE")
        return b"%PDF-invalid" if native == "invalid" else output
    monkeypatch.setattr(document_preview, "convert_word_to_pdf", convert)
    # A stale chunk must not override the fresh math-aware original fallback.
    from types import SimpleNamespace
    stale_chunks = AsyncMock(return_value=[SimpleNamespace(content="Stale text without equations")])
    monkeypatch.setattr(IngestionService, "list_document_chunks", stale_chunks)
    route = f"/api/v1/documents/{document_id}/preview"
    assert client.get(route).status_code == 401
    assert client.get(route, headers={"Authorization": f"Bearer test-token:{uuid.uuid4()}"}).status_code == 404
    response = client.get(route, headers=headers)
    assert response.status_code == 200
    assert response.headers["cache-control"] == "private, no-store"
    if native == "preserved":
        assert response.headers["content-type"] == "application/pdf"
        assert response.content == output
    else:
        assert response.headers["content-type"].startswith("text/plain")
        assert_equations(response.text)
    stale_chunks.assert_not_called()
    assert client.get(f"/api/v1/documents/{document_id}/download", headers=headers).content == equation_docx

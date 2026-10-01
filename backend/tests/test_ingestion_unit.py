"""Comprehensive unit tests for Phase 5 parsing, OCR abstraction, cleaning, and chunking."""
import io
import pytest
from PIL import Image
import fitz
import docx
import os
import shutil
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, ANY

from app.services.ingestion.cleaner import clean_text
from app.services.ingestion.models import DocumentPage, ProcessedChunk
from app.services.ingestion.ocr.base import OCRUnavailableError
from app.services.ingestion.ocr.mock import MockOCRProvider
from app.services.ingestion.parsers.pdf import PDFParser
from app.services.ingestion.parsers.docx import DocxParser
from app.services.ingestion.parsers.doc import LegacyDocParser
from app.services.ingestion.parsers.txt import TxtParser
from app.services.ingestion.parsers.image import ImageParser
from app.services.ingestion.parsers.base import EmptyDocumentError, DocumentCorruptError, UnsupportedDocumentError, ParserError
from app.services.ingestion.parsers.factory import get_parser_for_mime_or_filename
from app.services.ingestion.chunker import StructureAwareChunker
from app.services.ingestion.pipeline import IngestionPipeline
from app.services.ingestion.ocr.tesseract import TesseractOCRProvider


# ==============================================================================
# 1. Text Cleaner Tests
# ==============================================================================

def test_cleaner_deterministic_normalization():
    raw = "Hello   world!  This  is   a   test.\r\nLine two with\ttabs and spaces."
    res = clean_text(raw)
    assert res == "Hello world! This is a test.\nLine two with tabs and spaces."


def test_cleaner_dehyphenation():
    raw = "The infor-\nmation was pre-\nsented clearly across line breaks."
    res = clean_text(raw)
    assert "information" in res
    assert "presented" in res


def test_cleaner_strips_control_chars_but_preserves_newlines():
    raw = "Text with null \x00 and bell \x07 characters.\n\nNew paragraph."
    res = clean_text(raw)
    assert "\x00" not in res
    assert "\x07" not in res
    assert "New paragraph." in res


def test_cleaner_collapses_excessive_blank_lines():
    raw = "Paragraph 1\n\n\n\n\n\nParagraph 2\n\n\nParagraph 3"
    res = clean_text(raw)
    assert res == "Paragraph 1\n\nParagraph 2\n\nParagraph 3"


def test_cleaner_preserves_code_indentation():
    raw = "def solve():\n    x = 10\n    return x * 2"
    res = clean_text(raw)
    assert "def solve():\n    x = 10\n    return x * 2" == res


# ==============================================================================
# 2. PDF Parser Tests
# ==============================================================================

def _create_test_pdf(pages_text: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages_text:
        page = doc.new_page()
        if text:
            page.insert_text((50, 50), text)
    b = doc.tobytes()
    doc.close()
    return b


def _create_scanned_pdf(num_pages: int = 1) -> bytes:
    doc = fitz.open()
    img = Image.new("RGB", (100, 100), color=(200, 200, 200))
    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")
    raw_img = img_bytes.getvalue()

    for _ in range(num_pages):
        page = doc.new_page()
        rect = fitz.Rect(50, 50, 200, 200)
        page.insert_image(rect, stream=raw_img)
    b = doc.tobytes()
    doc.close()
    return b


def test_pdf_native_text_extraction_and_page_preservation():
    p1 = "Page 1: Overview of Distributed Systems and Fault Tolerance concepts in modern clusters."
    p2 = "Page 2: Consensus algorithms including Paxos and Raft protocols for state replication."
    pdf_bytes = _create_test_pdf([p1, p2])

    ocr = MockOCRProvider(available=False)  # OCR disabled; should not be needed
    parser = PDFParser(ocr_provider=ocr, min_text_chars=30)
    pages = parser.parse(pdf_bytes, "test_distributed.pdf")

    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert "Overview of Distributed Systems" in pages[0].text
    assert pages[0].used_ocr is False

    assert pages[1].page_number == 2
    assert "Consensus algorithms" in pages[1].text
    assert pages[1].used_ocr is False


def test_pdf_scanned_page_selective_ocr():
    scanned_bytes = _create_scanned_pdf(1)
    ocr = MockOCRProvider(available=True, return_text="Recognized OCR text from scanned exam sheet")
    parser = PDFParser(ocr_provider=ocr, min_text_chars=30)
    pages = parser.parse(scanned_bytes, "scanned.pdf")

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].used_ocr is True
    assert "Recognized OCR text from scanned exam sheet" in pages[0].text


def test_pdf_scanned_fails_when_ocr_unavailable():
    scanned_bytes = _create_scanned_pdf(1)
    ocr = MockOCRProvider(available=False)
    parser = PDFParser(ocr_provider=ocr, min_text_chars=30)

    with pytest.raises(OCRUnavailableError) as exc_info:
        parser.parse(scanned_bytes, "scanned_doc.pdf")
    assert "requiring OCR" in str(exc_info.value)



def test_pdf_mixed_document():
    # Page 1 has rich digital text; Page 2 is a scanned diagram/image
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "This is rich digital text for Page 1 of the hybrid document.")

    page2 = doc.new_page()
    img = Image.new("RGB", (100, 100), color=(150, 150, 150))
    img_b = io.BytesIO()
    img.save(img_b, format="PNG")
    page2.insert_image(fitz.Rect(50, 50, 200, 200), stream=img_b.getvalue())

    mixed_bytes = doc.tobytes()
    doc.close()

    ocr = MockOCRProvider(available=True, return_text="OCR text for scanned diagram on Page 2")
    parser = PDFParser(ocr_provider=ocr, min_text_chars=30)
    pages = parser.parse(mixed_bytes, "mixed.pdf")

    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert pages[0].used_ocr is False
    assert "rich digital text" in pages[0].text

    assert pages[1].page_number == 2
    assert pages[1].used_ocr is True
    assert "OCR text for scanned diagram" in pages[1].text


def test_pdf_corrupt_bytes_fails():
    parser = PDFParser(ocr_provider=MockOCRProvider())
    with pytest.raises(DocumentCorruptError):
        parser.parse(b"not a valid pdf", "broken.pdf")


# ==============================================================================
# 3. DOCX Parser Tests
# ==============================================================================

def _create_test_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Machine Learning Syllabus", level=1)
    doc.add_paragraph("Welcome to the advanced graduate syllabus.")
    doc.add_heading("Week 1 Topics", level=2)
    doc.add_paragraph("Supervised Learning", style="List Bullet")
    doc.add_paragraph("Unsupervised Learning", style="List Bullet")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Topic"
    table.cell(0, 1).text = "Hours"
    table.cell(1, 0).text = "Transformers"
    table.cell(1, 1).text = "10"

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_docx_structure_preservation():
    docx_bytes = _create_test_docx()
    parser = DocxParser()
    pages = parser.parse(docx_bytes, "syllabus.docx")

    assert len(pages) == 1
    text = pages[0].text

    # Check headings preserved with markdown markers
    assert "# Machine Learning Syllabus" in text
    assert "## Week 1 Topics" in text

    # Check list bullets preserved
    assert "- Supervised Learning" in text
    assert "- Unsupervised Learning" in text

    # Check table structure preserved
    assert "| Topic | Hours |" in text
    assert "| Transformers | 10 |" in text


def test_docx_empty_fails():
    doc = docx.Document()
    bio = io.BytesIO()
    doc.save(bio)
    parser = DocxParser()
    with pytest.raises(EmptyDocumentError):
        parser.parse(bio.getvalue(), "empty.docx")


# ==============================================================================
# 4. TXT Parser Tests
# ==============================================================================

def test_txt_parser_robust_decoding_and_structure():
    raw_content = "Heading: Physics 101\r\n\r\nNewton's first law explains inertia.\r\n\r\n    F = m * a"
    parser = TxtParser()
    pages = parser.parse(raw_content.encode("utf-8"), "physics.txt")

    assert len(pages) == 1
    assert "Newton's first law explains inertia." in pages[0].text
    assert "F = m * a" in pages[0].text


def test_txt_parser_latin1_encoding_fallback():
    # Text with Latin-1 specific character (e.g. copyright \xa9 or umlaut)
    content = "Tutor Copyright \xa9 2026. All rights reserved.".encode("latin-1")
    parser = TxtParser()
    pages = parser.parse(content, "legal.txt")
    assert len(pages) == 1
    assert "Tutor Copyright" in pages[0].text


def test_txt_parser_empty_fails():
    parser = TxtParser()
    with pytest.raises(EmptyDocumentError):
        parser.parse(b"   \n\n\t  ", "blank.txt")


# ==============================================================================
# 5. Image Parser Tests
# ==============================================================================

def test_image_parser_with_ocr():
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    bio = io.BytesIO()
    img.save(bio, format="PNG")

    ocr = MockOCRProvider(available=True, return_text="Diagram of Neural Network Architecture")
    parser = ImageParser(ocr_provider=ocr)
    pages = parser.parse(bio.getvalue(), "diagram.png")

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].used_ocr is True
    assert "Diagram of Neural Network Architecture" in pages[0].text


def test_image_parser_fails_when_ocr_unavailable():
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    bio = io.BytesIO()
    img.save(bio, format="PNG")

    ocr = MockOCRProvider(available=False)
    parser = ImageParser(ocr_provider=ocr)
    with pytest.raises(OCRUnavailableError):
        parser.parse(bio.getvalue(), "diagram.png")


# ==============================================================================
# 6. Parser Factory Tests
# ==============================================================================

def test_parser_factory_resolution():
    ocr = MockOCRProvider()
    assert isinstance(get_parser_for_mime_or_filename("application/pdf", "doc.pdf", ocr), PDFParser)
    assert isinstance(get_parser_for_mime_or_filename("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "doc.docx", ocr), DocxParser)
    assert isinstance(get_parser_for_mime_or_filename("application/msword", "doc.doc", ocr), LegacyDocParser)
    assert isinstance(get_parser_for_mime_or_filename("text/plain", "doc.txt", ocr), TxtParser)
    assert isinstance(get_parser_for_mime_or_filename("image/png", "doc.png", ocr), ImageParser)

    with pytest.raises(UnsupportedDocumentError):
        get_parser_for_mime_or_filename("application/zip", "archive.zip", ocr)


def _fake_tesseract(monkeypatch, languages):
    module = SimpleNamespace(
        pytesseract=SimpleNamespace(tesseract_cmd=None),
        get_languages=Mock(return_value=languages),
        image_to_string=Mock(return_value="recognized text"),
    )
    monkeypatch.setitem(sys.modules, "pytesseract", module)
    return module


def test_tesseract_provider_available_and_uses_configured_languages(monkeypatch):
    module = _fake_tesseract(monkeypatch, ["eng", "vie"])
    monkeypatch.setattr("app.services.ingestion.ocr.tesseract.shutil.which", lambda _: "/usr/bin/tesseract")
    provider = TesseractOCRProvider(languages="vie+eng")
    assert provider.is_available()
    assert provider.is_available()
    module.get_languages.assert_called_once_with(config="")
    assert provider.extract_text_from_image(object()) == "recognized text"
    module.image_to_string.assert_called_once_with(ANY, lang="vie+eng")


def test_tesseract_provider_unavailable_has_diagnostic(monkeypatch):
    monkeypatch.setitem(sys.modules, "pytesseract", SimpleNamespace())
    monkeypatch.setattr("app.services.ingestion.ocr.tesseract.shutil.which", lambda _: None)
    provider = TesseractOCRProvider()
    assert not provider.is_available()
    with pytest.raises(OCRUnavailableError, match="Tesseract executable is unavailable"):
        provider.extract_text_from_image(object())


def test_tesseract_provider_reports_missing_language_data(monkeypatch):
    _fake_tesseract(monkeypatch, ["eng"])
    monkeypatch.setattr("app.services.ingestion.ocr.tesseract.shutil.which", lambda _: "tesseract")
    provider = TesseractOCRProvider(languages="vie+eng")
    assert not provider.is_available()
    assert "vie" in provider.diagnostic


def test_tesseract_provider_uses_configured_executable(monkeypatch):
    module = _fake_tesseract(monkeypatch, ["eng", "vie"])
    monkeypatch.setattr("app.services.ingestion.ocr.tesseract.shutil.which", lambda _: None)
    provider = TesseractOCRProvider(tesseract_cmd="/opt/tools/tesseract", languages="eng")
    # An explicit path is checked directly and does not require PATH discovery.
    monkeypatch.setattr("app.services.ingestion.ocr.tesseract.os.path.isfile", lambda path: path == "/opt/tools/tesseract", raising=False)
    assert provider.is_available()
    assert module.pytesseract.tesseract_cmd == "/opt/tools/tesseract"


def test_legacy_doc_parser_mocked_conversion_and_cleanup(monkeypatch):
    converted_pages = [DocumentPage(1, "Converted content", "docx")]
    docx_parser = Mock(parse=Mock(return_value=converted_pages))
    parser = LegacyDocParser(libreoffice_cmd="soffice", docx_parser=docx_parser)
    parser.__dict__["executable"] = "soffice"
    observed = {}

    def fake_run(args, **kwargs):
        observed["args"] = args
        observed["kwargs"] = kwargs
        source = Path(args[-1])
        observed["source_bytes"] = source.read_bytes()
        observed["tempdir"] = args[args.index("--outdir") + 1]
        (Path(observed["tempdir"]) / "source.docx").write_bytes(b"converted")
        return subprocess.CompletedProcess(args, 0, "", "")

    monkeypatch.setattr("app.services.ingestion.parsers.doc.subprocess.run", fake_run)
    original = b"original legacy bytes"
    pages = parser.parse(original, "student notes.doc")
    assert pages == [DocumentPage(1, "Converted content", "doc")]
    assert observed["source_bytes"] == original
    assert observed["kwargs"]["timeout"] == parser.timeout_seconds
    assert observed["kwargs"].get("shell", False) is False
    assert not Path(observed["tempdir"]).exists()
    docx_parser.parse.assert_called_once_with(b"converted", "student notes.doc")


def test_legacy_doc_parser_missing_executable_is_clear(monkeypatch):
    parser = LegacyDocParser(libreoffice_cmd="missing-office")
    parser.__dict__["executable"] = None
    with pytest.raises(ParserError, match="requires LibreOffice"):
        parser.parse(b"doc", "notes.doc")


def test_legacy_doc_parser_timeout_is_reported(monkeypatch):
    parser = LegacyDocParser(timeout_seconds=1)
    parser.__dict__["executable"] = "soffice"
    monkeypatch.setattr(
        "app.services.ingestion.parsers.doc.subprocess.run",
        Mock(side_effect=subprocess.TimeoutExpired("soffice", 1)),
    )
    with pytest.raises(ParserError, match="exceeded the 1-second timeout"):
        parser.parse(b"doc", "notes.doc")


def test_legacy_doc_parser_conversion_failure_is_corrupt_document(monkeypatch):
    parser = LegacyDocParser()
    parser.__dict__["executable"] = "soffice"
    monkeypatch.setattr(
        "app.services.ingestion.parsers.doc.subprocess.run",
        Mock(return_value=subprocess.CompletedProcess(["soffice"], 1, "", "conversion failed")),
    )
    with pytest.raises(DocumentCorruptError, match="conversion failed"):
        parser.parse(b"corrupt doc", "broken.doc")


@pytest.mark.skipif(
    not (shutil.which("soffice") or os.getenv("LIBREOFFICE_CMD"))
    or not os.getenv("AI_TUTOR_LEGACY_DOC_FIXTURE"),
    reason="Set AI_TUTOR_LEGACY_DOC_FIXTURE and install/configure LibreOffice for the optional conversion smoke test.",
)
def test_legacy_doc_optional_environment_integration():
    from app.core.config import settings
    fixture = Path(os.environ["AI_TUTOR_LEGACY_DOC_FIXTURE"])
    assert fixture.is_file()
    pages = LegacyDocParser().parse(fixture.read_bytes(), fixture.name)
    assert pages
    assert pages[0].source_type == "doc"


# ==============================================================================
# 7. Structure-Aware Chunker Tests
# ==============================================================================

def test_chunker_basic_chunking_and_provenance():
    chunker = StructureAwareChunker(target_chunk_size=50, chunk_overlap=10, min_chunk_size=20)
    pages = [
        DocumentPage(page_number=1, text="This is page 1 with several sentences. It explains the main concept of machine learning algorithms.", source_type="pdf"),
        DocumentPage(page_number=2, text="This is page 2 detailing neural networks and deep learning backpropagation techniques.", source_type="pdf"),
    ]

    chunks = chunker.chunk_pages(pages)
    assert len(chunks) > 0

    # Verify sequential deterministic indices
    for idx, c in enumerate(chunks):
        assert c.chunk_index == idx
        assert c.token_count > 0
        assert c.page_number_start >= 1
        assert c.page_number_end >= c.page_number_start
        assert len(c.content.strip()) > 0


def test_chunker_multi_page_span_provenance():
    # Test that adjacent pages combined into a single chunk accurately track page_number_start and page_number_end
    chunker = StructureAwareChunker(target_chunk_size=100, chunk_overlap=0, min_chunk_size=10)
    pages = [
        DocumentPage(page_number=1, text="Short snippet on page 1.", source_type="pdf"),
        DocumentPage(page_number=2, text="Short snippet on page 2.", source_type="pdf"),
    ]

    chunks = chunker.chunk_pages(pages)
    assert len(chunks) == 1
    assert chunks[0].page_number_start == 1
    assert chunks[0].page_number_end == 2
    assert "page 1" in chunks[0].content
    assert "page 2" in chunks[0].content


def test_chunker_handles_very_long_paragraph():
    chunker = StructureAwareChunker(target_chunk_size=30, chunk_overlap=5, min_chunk_size=10)
    # Long paragraph without \n\n
    long_para = "Sentence one is here. " * 30
    pages = [DocumentPage(page_number=1, text=long_para, source_type="txt")]

    chunks = chunker.chunk_pages(pages)
    assert len(chunks) > 1
    for c in chunks:
        assert c.token_count <= 45  # Stays close to configured boundary


# ==============================================================================
# 8. Ingestion Pipeline End-to-End
# ==============================================================================

def test_ingestion_pipeline_pdf():
    pdf_bytes = _create_test_pdf([
        "Chapter 1: Principles of Operating Systems.\n\nVirtual memory provides an abstraction over physical RAM."
    ])
    pipeline = IngestionPipeline(ocr_provider=MockOCRProvider())
    chunks = pipeline.process(pdf_bytes, "os.pdf", "application/pdf")

    assert len(chunks) >= 1
    assert "Virtual memory provides an abstraction" in chunks[0].content
    assert chunks[0].page_number_start == 1
    assert chunks[0].page_number_end == 1
